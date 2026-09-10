import io
import sqlite3
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from database import get_db
except ImportError:
    from app.database import get_db

def _get_range_condition(range_filter: str):
    if range_filter == "today":
        return "DATE(tl.created_at) = DATE('now')"
    elif range_filter == "7days":
        return "DATE(tl.created_at) >= DATE('now', '-7 days')"
    elif range_filter == "month":
        return "STRFTIME('%Y-%m', tl.created_at) = STRFTIME('%Y-%m', 'now')"
    return "1=1"

def get_managerial_summary(area: str = "Todas", range_filter: str = "all") -> dict:
    conn = get_db()
    cur = conn.cursor()
    
    range_sql = _get_range_condition(range_filter)
    if area == "Todas":
        area_sql = "1=1"
        user_area_sql = "1=1"
        params = []
    elif area in ["Acceso", "Redes de Acceso"]:
        area_sql = "tl.area IN ('Soporte', 'Cabecera')"
        user_area_sql = "u.area IN ('Soporte', 'Cabecera')"
        params = []
    elif area in ["Servicios", "Servicios y Clientes"]:
        area_sql = "tl.area IN ('Telefonía')"
        user_area_sql = "u.area IN ('Telefonía')"
        params = []
    else:
        area_sql = "tl.area = ?"
        user_area_sql = "u.area = ?"
        params = [area]
    
    # 1. KPIs Globales
    cur.execute(f"""
    SELECT COUNT(tl.id) as total_tasks,
           COALESCE(SUM(tl.points), 0) as total_points,
           COALESCE(AVG(tl.net_duration), 0) as avg_mttr,
           COALESCE(SUM(tl.net_duration), 0) as total_net_min,
           COALESCE(SUM(tl.wait_minutes), 0) as total_wait_min
    FROM task_logs tl
    WHERE {range_sql} AND {area_sql}
    """, params)
    kpi_row = cur.fetchone()
    total_tasks = kpi_row[0] or 0
    total_points = kpi_row[1] or 0
    avg_mttr = round(kpi_row[2] or 0, 1)
    total_net_hours = round((kpi_row[3] or 0) / 60, 1)
    total_wait_hours = round((kpi_row[4] or 0) / 60, 1)
    
    # Cumplimiento SLA
    cur.execute(f"""
    SELECT COUNT(tl.id)
    FROM task_logs tl
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    WHERE {range_sql} AND {area_sql} AND tl.net_duration <= COALESCE(tt.sla_minutes, 45)
    """, params)
    within_sla = cur.fetchone()[0] or 0
    sla_compliance = round((within_sla / total_tasks * 100), 1) if total_tasks > 0 else 100.0

    # 2. Desglose por Células / Áreas Funcionales
    areas_list = ["Soporte", "Cabecera", "Telefonía"]
    area_breakdown = []
    for a in areas_list:
        cur.execute(f"""
        SELECT COUNT(tl.id), COALESCE(SUM(tl.points), 0), COALESCE(AVG(tl.net_duration), 0)
        FROM task_logs tl
        WHERE {range_sql} AND tl.area = ?
        """, [a])
        r = cur.fetchone()
        a_tasks, a_points, a_mttr = r[0] or 0, r[1] or 0, round(r[2] or 0, 1)
        
        # Cantidad de especialistas
        cur.execute("SELECT COUNT(*) FROM users WHERE area = ?", [a])
        techs_count = cur.fetchone()[0] or 4
        pts_per_tech = round(a_points / techs_count, 1) if techs_count else 0
        
        if pts_per_tech <= 25:
            status = "Equilibrada"
            badge = "success"
        elif pts_per_tech <= 45:
            status = "Moderada"
            badge = "warning"
        else:
            status = "Saturada / Alerta"
            badge = "danger"
            
        area_breakdown.append({
            "area": a,
            "total_tasks": a_tasks,
            "total_points": a_points,
            "avg_mttr": a_mttr,
            "techs_count": techs_count,
            "pts_per_tech": pts_per_tech,
            "share_percent": round((a_points / total_points * 100), 1) if total_points > 0 else 0,
            "status": status,
            "badge": badge
        })

    # 3. Productividad Detallada por Especialista (Dinámico)
    cur.execute(f"""
    SELECT u.id, u.name, u.area, u.role, u.avatar,
           COUNT(tl.id) as tasks_count,
           COALESCE(SUM(tl.points), 0) as total_points,
           COALESCE(AVG(tl.net_duration), 0) as avg_mttr,
           SUM(CASE WHEN tt.code LIKE 'P1%' THEN 1 ELSE 0 END) as p1_count,
           SUM(CASE WHEN tt.code LIKE 'P2%' THEN 1 ELSE 0 END) as p2_count,
           SUM(CASE WHEN tt.code LIKE 'P3%' THEN 1 ELSE 0 END) as p3_count,
           SUM(CASE WHEN tt.code LIKE 'P4%' THEN 1 ELSE 0 END) as p4_count,
           SUM(CASE WHEN tt.code LIKE 'P5%' THEN 1 ELSE 0 END) as p5_count
    FROM users u
    LEFT JOIN task_logs tl ON u.id = tl.user_id AND {range_sql}
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    WHERE {user_area_sql}
    GROUP BY u.id
    ORDER BY total_points DESC
    """, params)
    
    tech_rankings = []
    for row in cur.fetchall():
        pts = row[6] or 0
        if pts <= 25:
            st = "Equilibrada"
            st_color = "emerald"
        elif pts <= 45:
            st = "Moderada"
            st_color = "amber"
        else:
            st = "Alta"
            st_color = "red"
            
        tech_rankings.append({
            "id": row[0],
            "name": row[1],
            "area": row[2],
            "role": row[3],
            "avatar": row[4],
            "tasks_count": row[5] or 0,
            "total_points": pts,
            "avg_mttr": round(row[7] or 0, 1),
            "p1": row[8] or 0,
            "p2": row[9] or 0,
            "p3": row[10] or 0,
            "p4": row[11] or 0,
            "p5": row[12] or 0,
            "status": st,
            "status_color": st_color
        })

    # 4. Muestra de Registros de Auditoría
    cur.execute(f"""
    SELECT tl.ticket_code, tl.created_at, tl.duration_minutes, tl.wait_minutes, tl.net_duration,
           tl.points, tl.description, u.name as user_name, tl.area,
           et.subscriber_code, et.serial_pon, et.node_name, et.slot_pon, et.mac_address,
           tt.name as task_name, tt.code as task_code, tt.sla_minutes
    FROM task_logs tl
    LEFT JOIN users u ON tl.user_id = u.id
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    LEFT JOIN email_tickets et ON tl.ticket_code = et.ticket_code
    WHERE {range_sql} AND {area_sql}
    ORDER BY tl.id DESC
    LIMIT 30
    """, params)
    
    audit_samples = []
    for r in cur.fetchall():
        sub = r[9] or "N/A"
        perm = f"P-{sub[:2]}" if sub != "N/A" and len(sub) >= 2 else "N/A"
        audit_samples.append({
            "ticket_code": r[0],
            "date": r[1],
            "duration": r[2],
            "wait_time": r[3],
            "net_duration": r[4],
            "points": r[5],
            "notes": r[6] or "",
            "user_name": r[7],
            "area": r[8],
            "subscriber_code": sub,
            "permisor": perm,
            "serial_pon": r[10] or "N/A",
            "node_name": r[11] or "N/A",
            "slot_pon": r[12] or "N/A",
            "mac_address": r[13] or "N/A",
            "task_name": r[14] or "Operación",
            "task_code": r[15] or "P2",
            "sla_minutes": r[16] or 45
        })

    conn.close()
    return {
        "kpis": {
            "total_points": total_points,
            "total_tasks": total_tasks,
            "avg_mttr": avg_mttr,
            "total_net_hours": total_net_hours,
            "total_wait_hours": total_wait_hours,
            "sla_compliance": sla_compliance
        },
        "area_breakdown": area_breakdown,
        "tech_rankings": tech_rankings,
        "audit_samples": audit_samples,
        "filters": {
            "area": area,
            "range": range_filter
        }
    }


def generate_excel_report(area: str = "Todas", range_filter: str = "all") -> io.BytesIO:
    """
    Genera un libro Excel con 3 pestañas profesionales para supervisión gerencial:
    1. Resumen Ejecutivo y Áreas
    2. Productividad por Especialista
    3. Log Detallado de Auditoría
    """
    data = get_managerial_summary(area, range_filter)
    wb = openpyxl.Workbook()
    
    # Fuentes y estilos corporativos SnowUI / Inter
    f_title = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    f_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    f_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    f_bold = Font(name="Calibri", size=11, bold=True, color="1E293B")
    f_regular = Font(name="Calibri", size=10, color="1E293B")
    f_kpi_val = Font(name="Calibri", size=18, bold=True, color="1E3A8A")
    f_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="64748B")
    
    fill_navy = PatternFill("solid", fgColor="1E3A8A")
    fill_blue_head = PatternFill("solid", fgColor="2563EB")
    fill_ice = PatternFill("solid", fgColor="F8FAFC")
    fill_card = PatternFill("solid", fgColor="EFF6FF")
    
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================================
    # HOJA 1: RESUMEN EJECUTIVO Y ÁREAS
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo"
    ws1.views.sheetView[0].showGridLines = True
    
    # Encabezado Corporativo
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "INTER TELECOMUNICACIONES — REPORTE GERENCIAL OPERACIONES IP"
    ws1["A1"].font = f_title
    ws1["A1"].alignment = align_left
    
    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"División: Redes de Acceso y Aprovisionamiento | Filtro Área: {area} | Temporal: {range_filter.upper()} | Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A2"].font = f_subtitle
    ws1["A2"].alignment = align_left
    
    # 4 Tarjetas KPI Superiores
    kpis = data["kpis"]
    kpi_cards = [
        ("PUNTOS TOTALES", f"{kpis['total_points']} pts", "A4:B5"),
        ("TICKETS RESUELTOS", f"{kpis['total_tasks']}", "C4:D5"),
        ("MTTR PROMEDIO NETO", f"{kpis['avg_mttr']} min", "E4:E5"),
        ("CUMPLIMIENTO SLA", f"{kpis['sla_compliance']}%", "F4:G5")
    ]
    for lbl, val, range_str in kpi_cards:
        cells = list(ws1[range_str])
        top_left = cells[0][0]
        top_left.value = f"{lbl}\n{val}"
        top_left.font = f_bold
        top_left.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        top_left.fill = fill_card
        for row in cells:
            for c in row:
                c.border = border_thin
                c.fill = fill_card

    # Tabla: Balance de Células Funcionales
    ws1["A7"] = "BALANCE DE CARGA POR CÉLULA FUNCIONAL"
    ws1["A7"].font = f_bold
    
    headers_ws1 = ["Célula / Área", "Especialistas", "Tickets Resueltos", "Puntos Totales", "% Participación", "MTTR Promedio (min)", "Estado de Carga"]
    for col_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=8, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    curr_row = 9
    for ab in data["area_breakdown"]:
        ws1.cell(row=curr_row, column=1, value=ab["area"]).alignment = align_left
        ws1.cell(row=curr_row, column=2, value=ab["techs_count"]).alignment = align_center
        ws1.cell(row=curr_row, column=3, value=ab["total_tasks"]).alignment = align_right
        ws1.cell(row=curr_row, column=4, value=ab["total_points"]).alignment = align_right
        ws1.cell(row=curr_row, column=5, value=f"{ab['share_percent']}%").alignment = align_center
        ws1.cell(row=curr_row, column=6, value=ab["avg_mttr"]).alignment = align_right
        st_cell = ws1.cell(row=curr_row, column=7, value=ab["status"])
        st_cell.alignment = align_center
        if ab["status"] == "Equilibrada":
            st_cell.fill = PatternFill("solid", fgColor="D1FAE5")
        elif ab["status"] == "Moderada":
            st_cell.fill = PatternFill("solid", fgColor="FEF3C7")
        else:
            st_cell.fill = PatternFill("solid", fgColor="FEE2E2")
            
        for c in range(1, 8):
            ws1.cell(row=curr_row, column=c).border = border_thin
            ws1.cell(row=curr_row, column=c).font = f_regular
        curr_row += 1

    # =========================================================================
    # HOJA 2: PRODUCTIVIDAD POR ESPECIALISTA
    # =========================================================================
    ws2 = wb.create_sheet(title="Productividad Especialistas")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:M1")
    ws2["A1"] = "RENDIMIENTO Y DESGLOSE POR ESPECIALISTA"
    ws2["A1"].font = f_title
    
    headers_ws2 = [
        "Especialista", "Célula / Área", "Rol Funcional", "Tickets", "Puntos Totales", 
        "Pts/Ticket", "P1 (1pt)", "P2 (2pts)", "P3 (3pts)", "P4 (5pts)", "P5 (8pts)", 
        "MTTR Prom (min)", "Estado de Carga"
    ]
    for col_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_blue_head
        cell.alignment = align_center
        cell.border = border_thin
        
    row_idx = 4
    for tech in data["tech_rankings"]:
        ws2.cell(row=row_idx, column=1, value=tech["name"]).alignment = align_left
        ws2.cell(row=row_idx, column=2, value=tech["area"]).alignment = align_center
        ws2.cell(row=row_idx, column=3, value=tech["role"]).alignment = align_left
        ws2.cell(row=row_idx, column=4, value=tech["tasks_count"]).alignment = align_right
        
        pts_cell = ws2.cell(row=row_idx, column=5, value=tech["total_points"])
        pts_cell.alignment = align_right
        pts_cell.font = f_bold
        
        pts_per_t = round(tech["total_points"] / tech["tasks_count"], 1) if tech["tasks_count"] > 0 else 0
        ws2.cell(row=row_idx, column=6, value=pts_per_t).alignment = align_right
        
        ws2.cell(row=row_idx, column=7, value=tech["p1"]).alignment = align_center
        ws2.cell(row=row_idx, column=8, value=tech["p2"]).alignment = align_center
        ws2.cell(row=row_idx, column=9, value=tech["p3"]).alignment = align_center
        ws2.cell(row=row_idx, column=10, value=tech["p4"]).alignment = align_center
        ws2.cell(row=row_idx, column=11, value=tech["p5"]).alignment = align_center
        
        ws2.cell(row=row_idx, column=12, value=tech["avg_mttr"]).alignment = align_right
        
        st_cell = ws2.cell(row=row_idx, column=13, value=tech["status"])
        st_cell.alignment = align_center
        if tech["status"] == "Equilibrada":
            st_cell.fill = PatternFill("solid", fgColor="D1FAE5")
        elif tech["status"] == "Moderada":
            st_cell.fill = PatternFill("solid", fgColor="FEF3C7")
        else:
            st_cell.fill = PatternFill("solid", fgColor="FEE2E2")
            
        for c in range(1, 14):
            ws2.cell(row=row_idx, column=c).border = border_thin
            if c != 5:
                ws2.cell(row=row_idx, column=c).font = f_regular
        row_idx += 1

    # =========================================================================
    # HOJA 3: LOG DETALLADO DE AUDITORÍA
    # =========================================================================
    ws3 = wb.create_sheet(title="Log Forense Auditoría")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:R1")
    ws3["A1"] = "LOG DETALLADO DE AUDITORÍA Y TRAZABILIDAD TÉCNICA"
    ws3["A1"].font = f_title
    
    headers_ws3 = [
        "Ticket", "Fecha y Hora", "Duración Bruta (m)", "Pausa Terreno (m)", "Duración Neta (m)", 
        "Puntos", "Especialista", "Célula", "Abonado (10d)", "Permisor", "Serial PON (12c)", 
        "Nodo OLT", "Slot / PON", "MAC Abonado", "Tarea DERS", "Código Tarea", "SLA (min)", "Notas de Resolución"
    ]
    for col_idx, h in enumerate(headers_ws3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    # Obtener todos los registros para auditoría completa
    conn = get_db()
    cur = conn.cursor()
    range_sql = _get_range_condition(range_filter)
    if area == "Todas":
        area_sql = "1=1"
        params = []
    elif area in ["Acceso", "Redes de Acceso"]:
        area_sql = "tl.area IN ('Soporte', 'Cabecera')"
        params = []
    elif area in ["Servicios", "Servicios y Clientes"]:
        area_sql = "tl.area IN ('Telefonía')"
        params = []
    else:
        area_sql = "tl.area = ?"
        params = [area]
    
    cur.execute(f"""
    SELECT tl.ticket_code, tl.created_at, tl.duration_minutes, tl.wait_minutes, tl.net_duration,
           tl.points, u.name as user_name, tl.area,
           et.subscriber_code, et.serial_pon, et.node_name, et.slot_pon, et.mac_address,
           tt.name as task_name, tt.code as task_code, tt.sla_minutes, tl.description
    FROM task_logs tl
    LEFT JOIN users u ON tl.user_id = u.id
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    LEFT JOIN email_tickets et ON tl.ticket_code = et.ticket_code
    WHERE {range_sql} AND {area_sql}
    ORDER BY tl.id DESC
    """, params)
    
    audit_rows = cur.fetchall()
    conn.close()
    
    r_idx = 4
    for r in audit_rows:
        sub = r[8] or "N/A"
        perm = f"P-{sub[:2]}" if sub != "N/A" and len(sub) >= 2 else "N/A"
        
        ws3.cell(row=r_idx, column=1, value=r[0]).alignment = align_center
        ws3.cell(row=r_idx, column=2, value=r[1]).alignment = align_center
        ws3.cell(row=r_idx, column=3, value=r[2]).alignment = align_right
        ws3.cell(row=r_idx, column=4, value=r[3]).alignment = align_right
        ws3.cell(row=r_idx, column=5, value=r[4]).alignment = align_right
        ws3.cell(row=r_idx, column=6, value=r[5]).alignment = align_right
        ws3.cell(row=r_idx, column=7, value=r[6]).alignment = align_left
        ws3.cell(row=r_idx, column=8, value=r[7]).alignment = align_center
        ws3.cell(row=r_idx, column=9, value=sub).alignment = align_center
        ws3.cell(row=r_idx, column=10, value=perm).alignment = align_center
        ws3.cell(row=r_idx, column=11, value=r[9] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=12, value=r[10] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=13, value=r[11] or "N/A").alignment = align_left
        ws3.cell(row=r_idx, column=14, value=r[12] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=15, value=r[13] or "N/A").alignment = align_left
        ws3.cell(row=r_idx, column=16, value=r[14] or "P2").alignment = align_center
        ws3.cell(row=r_idx, column=17, value=r[15] or 45).alignment = align_right
        ws3.cell(row=r_idx, column=18, value=r[16] or "").alignment = align_left
        
        for c in range(1, 19):
            ws3.cell(row=r_idx, column=c).border = border_thin
            ws3.cell(row=r_idx, column=c).font = f_regular
        r_idx += 1

    # Autoajustar ancho de columnas en todas las hojas
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
